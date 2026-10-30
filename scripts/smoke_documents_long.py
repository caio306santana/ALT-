"""Real local validation of long attachments (PDF and DOCX) in the Notebook profile.

Synthetic files only, temporary database, real Ollama (qwen3.5:4b + EmbeddingGemma).
Checks what the owner is told about the part of the file the model saw, that a fact beyond
the excerpt is not invented, and that importing the file into a Project finds it.
Writes docs/documents-long-0.5.3-smoke.json, including failures."""
import io
import json
import sys
import tempfile
import time
import urllib.parse
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
from backend.api import create_app

MODEL='qwen3.5:4b'
START='O codigo do projeto e AURORA-17.'
END='O orcamento final aprovado foi de 48 mil reais.'
FILLER='Esta secao descreve rotinas de estudo, revisoes semanais e organizacao das tarefas da equipe.'

def long_pdf(pages=20):
    writer=PdfWriter()
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    for number in range(1,pages+1):
        page=writer.add_blank_page(width=612,height=792)
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
        lines=[f'Pagina {number}.']+([START] if number==1 else [])+[FILLER]*8+([END] if number==pages else [])
        body=''.join(f'BT /F1 9 Tf 40 {760-16*i} Td ({line}) Tj ET\n' for i,line in enumerate(lines))
        stream=DecodedStreamObject();stream.set_data(body.encode('ascii'));page[NameObject('/Contents')]=writer._add_object(stream)
    output=io.BytesIO();writer.write(output);return output.getvalue()

def long_docx(paragraphs=80):
    document=Document();document.add_heading('Plano do Projeto Aurora',1);document.add_paragraph(START)
    for i in range(paragraphs):document.add_paragraph(f'Paragrafo {i+1}. '+FILLER)
    document.add_paragraph(END)
    output=io.BytesIO();document.save(output);return output.getvalue()

token='synthetic-documents-smoke-'+'x'*40
report={'model':MODEL,'profile':'notebook','environment':'temporary database; synthetic files; real Ollama on this notebook','cases':{}}
output=Path(__file__).resolve().parents[1]/'docs'/'documents-long-0.5.3-smoke.json'
def save():output.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def chat(client,cid,content,attachment_ids=()):
    started=time.monotonic()
    response=client.post('/chat',json={'conversation_id':cid,'model':MODEL,'content':content,'attachment_ids':list(attachment_ids)})
    events=[json.loads(line) for line in response.text.splitlines()] if response.status_code==200 else []
    start=next((e for e in events if e['type']=='start'),{})
    return {'http':response.status_code,'seconds':round(time.monotonic()-started,1),'status':events[-1]['status'] if events else response.text[:300],
        'reply':''.join(e.get('text','') for e in events if e['type']=='delta'),'documents_notice':start.get('documents'),
        'recalled':start.get('recalled_memories'),'errors':[e.get('message') for e in events if e['type']=='error']}

with tempfile.TemporaryDirectory(prefix='alt-documents-smoke-') as directory:
    data=Path(directory)
    with TestClient(create_app(data,token),headers={'Authorization':'Bearer '+token}) as client:
        prefs=client.get('/settings').json();prefs['performance']='notebook';client.put('/settings',json=prefs)
        for name,content in (('plano-longo.pdf',long_pdf()),('plano-longo.docx',long_docx())):
            attached=client.post('/attachments',content=content,headers={'x-alt-name':urllib.parse.quote(name)}).json()
            case={'file':name,'bytes':len(content),'extracted_chars':attached.get('chars'),'notice':attached.get('notice')}
            cid=client.post('/conversations').json()['id']
            case['end_fact']=chat(client,cid,'Qual foi o orçamento final aprovado, segundo o anexo? Se o trecho que você recebeu não mostrar, diga isso.',[attached['id']])
            case['end_fact']['invented_value']=any(v in case['end_fact']['reply'] for v in ('48','quarenta e oito')) is False and any(ch.isdigit() for ch in case['end_fact']['reply'].replace('17',''))
            case['end_fact']['answered_correctly']='48' in case['end_fact']['reply']
            case['end_fact']['says_only_beginning']=any(w in case['end_fact']['reply'].casefold() for w in ('início','inicio','trecho','parte'))
            # An attachment belongs to the conversation it was sent in: upload again for a new one.
            again=client.post('/attachments',content=content,headers={'x-alt-name':urllib.parse.quote(name)}).json()
            cid=client.post('/conversations').json()['id']
            case['start_fact']=chat(client,cid,'Qual é o código do projeto no anexo? Responda só o código.',[again['id']])
            case['start_fact']['correct']='AURORA-17' in case['start_fact']['reply'].upper()
            report['cases'][name]=case;save();print(json.dumps({name:case},ensure_ascii=False)[:1200],flush=True)
        # Whole document through a Project: semantic search reaches the last page.
        pdf=data/'plano-longo.pdf';pdf.write_bytes(long_pdf())
        project=client.post('/projects',json={'name':'Aurora'}).json()
        imported=client.post(f"/projects/{project['id']}/documents",json={'path':str(pdf)}).json()
        client.post('/knowledge/index')
        deadline=time.monotonic()+900
        while time.monotonic()<deadline and client.get('/knowledge/status').json()['status']=='running':time.sleep(2)
        hits=client.post('/knowledge/search',json={'query':'orçamento final aprovado'}).json()
        policy=client.get('/policy').json()['policy'];policy['memory_recall']='semantic';client.put('/policy',json=policy)
        cid=client.post('/conversations').json()['id'];client.put(f'/conversations/{cid}/project',json={'project_id':project['id']})
        answer=chat(client,cid,'Qual foi o orçamento final aprovado do Projeto Aurora?')
        answer['answered_correctly']='48' in answer['reply']
        report['project']={'chunks':imported.get('chunks'),'index':client.get('/knowledge/status').json(),
            'top_hit_has_end_fact':bool(hits) and '48 mil' in json.dumps(hits[:2],ensure_ascii=False),'chat':answer}
        save();print(json.dumps({'project':report['project']},ensure_ascii=False)[:1500],flush=True)
