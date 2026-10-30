"""Exercise synthetic attachments through the real local API and installed Ollama models."""
import argparse
import io
import json
import sys
import tempfile
import time
import urllib.parse
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
from backend.api import create_app


def native_pdf(text):
    writer=PdfWriter();page=writer.add_blank_page(width=612,height=792)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    stream=DecodedStreamObject();stream.set_data(('BT /F1 20 Tf 60 700 Td ('+text+') Tj ET').encode('ascii'))
    page[NameObject('/Contents')]=writer._add_object(stream)
    output=io.BytesIO();writer.write(output);return output.getvalue()


def fixtures():
    word=Document();word.add_paragraph('Validacao de documento')
    table=word.add_table(rows=1,cols=2);table.cell(0,0).text='Codigo';table.cell(0,1).text='ALT-733'
    output=io.BytesIO();word.save(output)
    image=Image.new('RGB',(1240,1754),'white')
    draw=ImageDraw.Draw(image);font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',54)
    draw.text((100,120),'CODIGO: ALT-734',font=font,fill='black')
    draw.text((100,240),'Reuniao: 15 horas',font=font,fill='black')
    scan=io.BytesIO();image.save(scan,format='PDF',resolution=150)
    jpeg=io.BytesIO();image.save(jpeg,format='JPEG')
    return [('note.txt',b'Codigo de teste: ALT-731','ALT-731'),
            ('native.pdf',native_pdf('Codigo de teste: ALT-732'),'ALT-732'),
            ('table.docx',output.getvalue(),'ALT-733'),
            ('scan.pdf',scan.getvalue(),'ALT-734'),
            ('image.jpg',jpeg.getvalue(),'ALT-734')]


def main(args):
    root=Path(__file__).resolve().parents[1]
    (root/'.local').mkdir(exist_ok=True)
    token='synthetic-document-smoke-'+'x'*40
    report={'synthetic_only':True,'cloud_used':False,'results':[]}
    output=Path(args.output);output.parent.mkdir(parents=True,exist_ok=True)
    def save():output.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    with tempfile.TemporaryDirectory(prefix='document-smoke-',dir=root/'.local') as directory:
        app=create_app(Path(directory),token)
        with TestClient(app,headers={'Authorization':'Bearer '+token}) as client:
            installed={row['name'] for row in client.get('/models').json()['models']}
            if set(args.models)-installed:raise SystemExit('Somente modelos locais já instalados são permitidos. Nenhum download será feito.')
            prefs=client.get('/settings').json();prefs['performance']='desktop';client.put('/settings',json=prefs)
            for model in args.models:
                for name,data,expected in fixtures():
                    row={'model':model,'file':name,'expected':expected};start=time.monotonic()
                    uploaded=client.post('/attachments',content=data,headers={'X-Alt-Name':urllib.parse.quote(name),'Content-Type':'application/octet-stream'})
                    row['upload_status']=uploaded.status_code
                    if uploaded.status_code==200:
                        attachment=uploaded.json();row['notice']=attachment.get('notice')
                        cid=client.post('/conversations').json()['id']
                        response=client.post('/chat',json={'conversation_id':cid,'model':model,'content':'Qual e o codigo no anexo? Responda somente com o codigo.','attachment_ids':[attachment['id']]})
                        row['chat_status']=response.status_code
                        if response.status_code==200:
                            events=[json.loads(line) for line in response.text.splitlines()]
                            row['reply']=''.join(event.get('text','') for event in events if event['type']=='delta')
                            row['metrics']=events[-1].get('metrics')
                            row['ok']=events[-1].get('status')=='complete' and expected in row['reply']
                        else:row['error']=response.json()
                    else:row['error']=uploaded.json()
                    row['seconds']=round(time.monotonic()-start,3)
                    report['results'].append(row);save();print(json.dumps(row,ensure_ascii=True),flush=True)
    report['passed']=sum(bool(row.get('ok')) for row in report['results']);report['total']=len(report['results']);save()
    return 0 if report['passed']==report['total'] else 1


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--models',nargs='+',default=['qwen3.5:4b','qwen3.5:9b'])
    parser.add_argument('--output',default='.local/document-chat-smoke.json')
    raise SystemExit(main(parser.parse_args()))
