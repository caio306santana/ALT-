"""Local project documents and genuine vector retrieval through Ollama embeds."""
import asyncio
import hashlib
import json
import math
import stat
import zipfile
from pathlib import Path
from uuid import uuid4
from backend.memory.store import now
from backend.memory.library import MemoryBody,reject_secrets

def unit(vector):
    if not vector or len(vector)>4096 or any(not isinstance(x,(float,int)) or not math.isfinite(x) for x in vector):raise ValueError('Vetor inválido do modelo local.')
    norm=math.sqrt(sum(x*x for x in vector))
    if not norm:raise ValueError('Vetor vazio do modelo local.')
    return [x/norm for x in vector]

def chunks(text,size=1100,overlap=100):
    text=text.strip()
    for start in range(0,len(text),size-overlap):
        part=text[start:start+size]
        if part.strip():yield part

def extract_document(path):
    path=Path(path)
    if not path.is_absolute() or str(path).startswith(('\\\\','//')) or path.suffix.lower() not in ('.pdf','.docx','.txt','.md'):raise ValueError('Escolha PDF, DOCX, TXT ou Markdown local.')
    for parent in [path,*path.parents]:
        info=parent.lstat()
        if stat.S_ISLNK(info.st_mode) or getattr(info,'st_file_attributes',0)&0x400:raise ValueError('Importação não aceita links ou junctions.')
    if not path.is_file() or path.stat().st_size>10_000_000:raise ValueError('Documento deve ter até 10 MB.')
    sections=[]
    if path.suffix.lower()=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted:raise ValueError('PDF protegido; exporte uma cópia de texto autorizada para importar.')
        if len(reader.pages)>200:raise ValueError('PDF excede 200 páginas.')
        for index,page in enumerate(reader.pages):sections.append((f'página {index+1}',page.extract_text() or ''))
    elif path.suffix.lower()=='.docx':
        from docx import Document
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist())>50_000_000:raise ValueError('DOCX expandido excede o limite.')
        document=Document(path)
        sections=[(f'parágrafo {i+1}',p.text) for i,p in enumerate(document.paragraphs) if p.text.strip()]
        for index,table in enumerate(document.tables):sections.append((f'tabela {index+1}','\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows)))
    else:
        raw=path.read_bytes()
        if b'\x00' in raw:raise ValueError('Arquivo não é texto UTF-8/Windows-1252.')
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:text=raw.decode('cp1252')
        sections=[('texto',text)]
    text='\n'.join(t for _,t in sections)
    if not text.strip():raise ValueError('Não há texto extraível. PDFs digitalizados precisam de OCR, ainda não incluído.')
    if len(text)>200_000:raise ValueError('Texto extraído excede 200.000 caracteres.')
    reject_secrets(text)
    return sections

class Knowledge:
    def __init__(self,store,library,continuity,engine,ollama):
        self.store,self.library,self.continuity,self.engine,self.ollama=store,library,continuity,engine,ollama
        self.index_task=None;self.summary_tasks={};self.query_vectors={};self.index_state={'status':'idle','completed':0,'total':0}
        store.db.executescript('''
          CREATE TABLE IF NOT EXISTS projects(id TEXT PRIMARY KEY,name TEXT NOT NULL,created TEXT NOT NULL);
          CREATE TABLE IF NOT EXISTS conversation_projects(conversation_id TEXT PRIMARY KEY REFERENCES conversations(id) ON DELETE CASCADE,project_id TEXT REFERENCES projects(id) ON DELETE CASCADE,last_summary_count INTEGER NOT NULL DEFAULT 0);
          CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY,project_id TEXT REFERENCES projects(id) ON DELETE CASCADE,name TEXT NOT NULL,sha256 TEXT NOT NULL,created TEXT NOT NULL);
          CREATE TABLE IF NOT EXISTS document_chunks(id TEXT PRIMARY KEY,document_id TEXT REFERENCES documents(id) ON DELETE CASCADE,location TEXT NOT NULL,content TEXT NOT NULL);
          CREATE TABLE IF NOT EXISTS semantic_vectors(kind TEXT NOT NULL,item_id TEXT NOT NULL,content_hash TEXT NOT NULL,model_digest TEXT NOT NULL,vector TEXT NOT NULL,PRIMARY KEY(kind,item_id));
          CREATE TABLE IF NOT EXISTS project_memories(project_id TEXT REFERENCES projects(id) ON DELETE CASCADE,memory_id TEXT REFERENCES memories(id) ON DELETE CASCADE,PRIMARY KEY(project_id,memory_id));
          CREATE TABLE IF NOT EXISTS project_summary_drafts(id TEXT PRIMARY KEY,project_id TEXT REFERENCES projects(id) ON DELETE CASCADE,conversation_id TEXT REFERENCES conversations(id) ON DELETE CASCADE,payload TEXT NOT NULL,status TEXT NOT NULL,created TEXT NOT NULL);
        ''');store.db.commit()

    def projects(self):return self.store.rows('SELECT * FROM projects ORDER BY created DESC')
    def create_project(self,name):
        name=name.strip();reject_secrets(name)
        if not name or len(name)>100:raise ValueError('Nome do projeto deve ter de 1 a 100 caracteres.')
        if len(self.projects())>=100:raise ValueError('Limite de 100 projetos.')
        pid=str(uuid4());self.store.write('INSERT INTO projects VALUES(?,?,?)',(pid,name,now()));return {'id':pid,'name':name}
    def link(self,cid,pid):
        if not self.store.conversation(cid):raise ValueError('Conversa não encontrada.')
        if pid and not self.store.rows('SELECT id FROM projects WHERE id=?',(pid,)):raise ValueError('Projeto não encontrado.')
        if pid:self.store.write('INSERT INTO conversation_projects VALUES(?,?,0) ON CONFLICT(conversation_id) DO UPDATE SET project_id=excluded.project_id,last_summary_count=0',(cid,pid))
        else:self.store.write('DELETE FROM conversation_projects WHERE conversation_id=?',(cid,))
        return {'project_id':pid}
    def project_for(self,cid):
        rows=self.store.rows('SELECT project_id FROM conversation_projects WHERE conversation_id=?',(cid,));return rows[0]['project_id'] if rows else ''
    def documents(self,pid):return self.store.rows('SELECT d.*,(SELECT COUNT(*) FROM document_chunks c WHERE c.document_id=d.id) AS chunks FROM documents d WHERE d.project_id=? ORDER BY d.created DESC',(pid,))
    def import_document(self,pid,path):
        if not self.store.rows('SELECT id FROM projects WHERE id=?',(pid,)):raise ValueError('Projeto não encontrado.')
        sections=extract_document(path);all_text='\n'.join(t for _,t in sections);digest=hashlib.sha256(all_text.encode()).hexdigest()
        duplicate=self.store.rows('SELECT id FROM documents WHERE project_id=? AND sha256=?',(pid,digest))
        if duplicate:return {'id':duplicate[0]['id'],'duplicate':True}
        if len(self.documents(pid))>=100:raise ValueError('Limite de 100 documentos por projeto.')
        did=str(uuid4());rows=[(str(uuid4()),did,loc,text) for loc,section in sections for text in chunks(section)]
        with self.store.lock:
            try:
                self.store.db.execute('INSERT INTO documents VALUES(?,?,?,?,?)',(did,pid,Path(path).name,digest,now()))
                self.store.db.executemany('INSERT INTO document_chunks VALUES(?,?,?,?)',rows);self.store.db.commit()
            except Exception:self.store.db.rollback();raise
        self.store.audit('document_imported',did);return {'id':did,'chunks':len(rows),'name':Path(path).name,'duplicate':False}

    async def model_digest(self):
        model=next((m for m in await self.ollama.models() if m.get('name','').split(':')[0]=='embeddinggemma' and not m.get('remote_host')),None)
        if not model:raise ValueError('EmbeddingGemma não está instalado. Autorize o download para ativar busca semântica.')
        return model.get('digest') or model['name']
    def items(self,pid=''):
        memories=self.library.list();rows=[]
        for row in memories:
            if not row['enabled']:continue
            projects=self.store.rows('SELECT project_id FROM project_memories WHERE memory_id=?',(row['id'],))
            if pid and projects and not any(p['project_id']==pid for p in projects):continue
            rows.append({'kind':'memory','id':row['id'],'title':row['title'],'content':row['content'],'category':row['category']})
        query='SELECT c.*,d.name,d.project_id FROM document_chunks c JOIN documents d ON d.id=c.document_id'
        docs=self.store.rows(query+(' WHERE d.project_id=?' if pid else ''),(pid,) if pid else ())
        rows.extend({'kind':'chunk','id':d['id'],'title':d['name']+' · '+d['location'],'content':d['content'],'category':'contexto','document_id':d['document_id']} for d in docs)
        return rows
    async def index(self):
        if self.index_task and not self.index_task.done():raise ValueError('Indexação já está em andamento.')
        digest=await self.model_digest();cancel=self.engine.begin('semantic-index')
        self.index_task=asyncio.create_task(self._index(digest,cancel));return {'started':True}
    async def _index(self,digest,cancel):
        self.index_state={'status':'running','completed':0,'total':0}
        try:
            items=self.items();self.index_state['total']=len(items)
            for item in items:
                if cancel.is_set():break
                text=item['title']+'\n'+item['content'];fingerprint=hashlib.sha256(text.encode()).hexdigest()
                existing=self.store.rows('SELECT * FROM semantic_vectors WHERE kind=? AND item_id=?',(item['kind'],item['id']))
                if not existing or existing[0]['content_hash']!=fingerprint or existing[0]['model_digest']!=digest:
                    # Long memories are encoded in chunks and averaged, not silently
                    # truncated beyond the embedding model's input window.
                    segments=list(chunks(text));embedding=asyncio.create_task(self.ollama.embed(segments));stopped=asyncio.create_task(cancel.wait())
                    try:
                        ready,_=await asyncio.wait([embedding,stopped],return_when=asyncio.FIRST_COMPLETED)
                        if stopped in ready:break
                        values=[unit(v) for v in embedding.result()]
                        if not values or len({len(v) for v in values})!=1:raise ValueError('Dimensões inconsistentes de embeddings.')
                        vector=unit([sum(v[i] for v in values)/len(values) for i in range(len(values[0]))])
                    finally:
                        for task in (embedding,stopped):
                            if not task.done():task.cancel()
                        await asyncio.gather(embedding,stopped,return_exceptions=True)
                    self.store.write('INSERT INTO semantic_vectors VALUES(?,?,?,?,?) ON CONFLICT(kind,item_id) DO UPDATE SET content_hash=excluded.content_hash,model_digest=excluded.model_digest,vector=excluded.vector',(item['kind'],item['id'],fingerprint,digest,json.dumps(vector)))
                self.index_state['completed']+=1
            self.index_state['status']='cancelled' if cancel.is_set() else 'complete'
        except Exception as exc:self.index_state.update({'status':'error','error':str(exc)[:400]})
        finally:self.engine.active.pop('semantic-index',None)
    async def search(self,query,pid='',limit=6,minimum=.25):
        query=query.strip()
        if not query:return []
        digest=await self.model_digest();key=(digest,query[:1500])
        if key not in self.query_vectors:
            # Typing re-estimates context; repeated queries reuse their vector.
            if len(self.query_vectors)>=64:self.query_vectors.pop(next(iter(self.query_vectors)))
            self.query_vectors[key]=unit((await self.ollama.embed([query[:1500]]))[0])
        vector=self.query_vectors[key];ranked=[]
        stored={(r['kind'],r['item_id']):r for r in self.store.rows('SELECT * FROM semantic_vectors WHERE model_digest=?',(digest,))}
        for item in self.items(pid):
            row=stored.get((item['kind'],item['id']))
            # Vectors of edited items are stale until re-indexed; never rank them.
            if not row or row['content_hash']!=hashlib.sha256((item['title']+'\n'+item['content']).encode()).hexdigest():continue
            indexed=json.loads(row['vector'])
            if len(indexed)!=len(vector):continue
            score=sum(a*b for a,b in zip(vector,indexed))
            if score>=minimum:ranked.append({**item,'score':round(score,4)})
        ranked.sort(key=lambda r:r['score'],reverse=True);return ranked[:limit]
    def status(self):
        indexed=self.store.rows('SELECT COUNT(*) AS n FROM semantic_vectors')[0]['n']
        return {**self.index_state,'indexed':indexed,'items':len(self.items())}
    def delete_document(self,did):
        with self.store.lock:
            chunk_ids=[r['id'] for r in self.store.rows('SELECT id FROM document_chunks WHERE document_id=?',(did,))]
            self.store.db.executemany("DELETE FROM semantic_vectors WHERE kind='chunk' AND item_id=?",[(c,) for c in chunk_ids])
            self.store.db.execute('DELETE FROM document_chunks WHERE document_id=?',(did,));self.store.db.execute('DELETE FROM documents WHERE id=?',(did,));self.store.db.commit()
        self.store.audit('document_deleted',did);return {'ok':True}
    def delete_project(self,pid):
        for doc in self.documents(pid):self.delete_document(doc['id'])
        # Approved memories survive; they only lose the project link.
        self.store.write('DELETE FROM projects WHERE id=?',(pid,));self.store.audit('project_deleted',pid);return {'ok':True}

    def summary_drafts(self,pid):return [{**r,'payload':json.loads(r['payload'])} for r in self.store.rows("SELECT * FROM project_summary_drafts WHERE project_id=? AND status='draft' ORDER BY created DESC",(pid,))]
    async def draft_summary(self,cid,model):
        pid=self.project_for(cid)
        if not pid:raise ValueError('Vincule a conversa a um projeto primeiro.')
        draft=await self.continuity.draft(cid,model);internal=self.continuity.drafts[draft['id']]
        payload={**draft,'fingerprint':internal['fingerprint'],'message_ids':internal['message_ids']}
        self.store.write('INSERT INTO project_summary_drafts VALUES(?,?,?,?,?,?)',(draft['id'],pid,cid,json.dumps(payload,ensure_ascii=False),'draft',now()))
        count=len(self.continuity.snapshot(cid)[1]);self.store.write('UPDATE conversation_projects SET last_summary_count=? WHERE conversation_id=?',(count,cid))
        return draft
    def approve_summary(self,did,body):
        rows=self.store.rows("SELECT * FROM project_summary_drafts WHERE id=? AND status='draft'",(did,))
        if not rows:raise ValueError('Rascunho não encontrado.')
        row=rows[0];payload=json.loads(row['payload'])
        if self.continuity.snapshot(row['conversation_id'])[2]!=payload['fingerprint']:raise ValueError('A conversa mudou; gere um resumo novo.')
        memory=self.library.save(body)
        self.store.write('INSERT INTO memory_sources VALUES(?,?,?)',(memory['id'],row['conversation_id'],json.dumps(payload['message_ids'])))
        self.store.write('INSERT INTO project_memories VALUES(?,?)',(row['project_id'],memory['id']))
        self.store.write("UPDATE project_summary_drafts SET status='approved' WHERE id=?",(did,));return memory
    def schedule_summary(self,cid,model):
        if model.startswith('cloud:') or not self.store.settings().get('owner_policy',{}).get('auto_project_summaries'):return
        rows=self.store.rows('SELECT * FROM conversation_projects WHERE conversation_id=?',(cid,))
        if not rows or len(self.continuity.snapshot(cid)[1])-rows[0]['last_summary_count']<12:return
        if cid in self.summary_tasks and not self.summary_tasks[cid].done():return
        async def generate():
            try:await self.draft_summary(cid,model)
            except Exception:self.store.audit('project_summary_not_generated',cid)
        self.summary_tasks[cid]=asyncio.create_task(generate())
