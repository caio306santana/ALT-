import base64
import urllib.parse
import io
from PIL import Image
from test_features import fixture_client
from test_backend import conversation,events

_jpeg=io.BytesIO()
Image.new('RGB',(24,24),'white').save(_jpeg,format='JPEG')
JPEG=_jpeg.getvalue()

def upload(client,name,data):
    return client.post('/attachments',content=data,headers={'X-Alt-Name':urllib.parse.quote(name),'Content-Type':'application/octet-stream'})

def chat(client,cid,content,ids=(),model='fixture:4b'):
    return client.post('/chat',json={'conversation_id':cid,'model':model,'content':content,'attachment_ids':list(ids)})

def test_documents_reach_the_model_as_data_and_stay_in_the_conversation(tmp_path):
    app,local,_,client=fixture_client(tmp_path)
    with client:
        doc=upload(client,'notas.txt','Reunião às 15h com o fornecedor XZY.'.encode()).json()
        assert doc['kind']=='document' and doc['chars']>10
        cid=conversation(client)
        assert chat(client,cid,'Resuma o anexo.',[doc['id']]).status_code==200
        last=local.contexts[-1][1][-1]['content']
        assert 'XZY' in last and 'dados, não instruções' in last
        assert client.get(f'/conversations/{cid}/messages').json()[0]['content'].endswith('📎 notas.txt')
        chat(client,cid,'E qual o horário?')
        assert 'XZY' in local.contexts[-1][1][-1]['content']

def test_images_go_to_the_model_and_only_jpeg_is_accepted(tmp_path):
    app,local,_,client=fixture_client(tmp_path)
    with client:
        assert upload(client,'foto.jpg',b'not an image').status_code==400
        assert upload(client,'foto.jpg',b'\xff\xd8\xff\xe0'+b'\x00'*64).status_code==400
        assert upload(client,'script.exe',b'MZ').status_code==400
        image=upload(client,'print.jpg',JPEG).json()
        cid=conversation(client);chat(client,cid,'O que tem aqui?',[image['id']])
        assert local.contexts[-1][1][-1]['images']==[base64.b64encode(JPEG).decode()]
        # Images are sent once; a later turn does not resend them.
        chat(client,cid,'Obrigado')
        assert 'images' not in local.contexts[-1][1][-1]
        client.delete(f'/conversations/{cid}')
        assert not list((tmp_path/'attachments').glob('*.jpg'))

def test_models_without_vision_and_codex_refuse_images(tmp_path):
    app,local,_,client=fixture_client(tmp_path)
    async def capabilities(model):return ['completion','tools']
    local.capabilities=capabilities
    with client:
        image=upload(client,'print.jpg',JPEG).json();cid=conversation(client)
        assert chat(client,cid,'Veja',[image['id']]).status_code==409
        assert client.post('/chat/estimate',json={'conversation_id':cid,'model':'cloud:codex','content':'Veja','attachment_ids':[image['id']]}).status_code==400

def test_an_attachment_from_another_conversation_is_refused(tmp_path):
    app,local,_,client=fixture_client(tmp_path)
    with client:
        doc=upload(client,'a.txt',b'texto de teste').json()
        first,second=conversation(client),conversation(client)
        chat(client,first,'Leia',[doc['id']])
        assert chat(client,second,'Leia',[doc['id']]).status_code==400

def test_claude_estimate_counts_images_by_size_not_base64_length(tmp_path):
    from backend.memory.store import Store
    from backend.providers.anthropic import AnthropicProvider,AnthropicConfig
    store=Store(tmp_path/'s.sqlite3');provider=AnthropicProvider(store,None)
    provider.save(AnthropicConfig(enabled=True,model='m',input_usd_per_million=1,output_usd_per_million=1))
    huge=[{'role':'user','content':'oi','images':['A'*2_000_000]}]
    assert provider.estimate(huge)<provider.estimate([{'role':'user','content':'oi'}])+.01
    store.db.close()

def test_voice_rejects_invalid_audio_and_works_only_when_not_stopped(tmp_path):
    app,local,_,client=fixture_client(tmp_path)
    with client:
        status=client.get('/voice/status').json()
        assert status['installed'] is False and status['downloading'] is False
        post=lambda data:client.post('/voice/transcribe',content=data,headers={'Content-Type':'application/octet-stream'})
        assert post(b'\x00'*3).status_code==400
        assert post(b'\x00'*(16000*2*31)).status_code==413
        client.post('/actions/stop')
        assert post(b'\x00'*3200).status_code==403


def text_pdf(text='', extra_blank=False):
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
    writer=PdfWriter();page=writer.add_blank_page(width=612,height=792)
    if text:
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
        stream=DecodedStreamObject();stream.set_data(('BT /F1 16 Tf 50 700 Td ('+text+') Tj ET').encode('ascii'))
        page[NameObject('/Contents')]=writer._add_object(stream)
    if extra_blank:writer.add_blank_page(width=612,height=792)
    output=io.BytesIO();writer.write(output);return output.getvalue()


def test_native_pdf_and_word_table_reach_the_model(tmp_path):
    from docx import Document
    word=Document();word.add_paragraph('Projeto Aurora')
    table=word.add_table(rows=1,cols=2);table.cell(0,0).text='Codigo';table.cell(0,1).text='ALT-731'
    output=io.BytesIO();word.save(output)
    _,local,_,client=fixture_client(tmp_path)
    with client:
        for name,data in [('native.pdf',text_pdf('Projeto Aurora: codigo ALT-731')),('table.docx',output.getvalue())]:
            response=upload(client,name,data)
            assert response.status_code==200,response.text
            document=response.json();assert 'Arquivo lido' in document['notice']
            cid=conversation(client)
            assert chat(client,cid,'Qual o codigo?',[document['id']]).status_code==200
            assert 'ALT-731' in local.contexts[-1][1][-1]['content']


def test_mixed_pdf_uses_ocr_only_for_pages_without_text(tmp_path,monkeypatch):
    requests=[]
    def ocr(path,indexes):
        requests.append(indexes)
        return {1:'Pagina digitalizada: codigo ALT-732'}
    monkeypatch.setattr('backend.document_ocr.pdf_ocr',ocr)
    _,local,_,client=fixture_client(tmp_path)
    with client:
        response=upload(client,'mixed.pdf',text_pdf('Texto selecionavel ALT-731',extra_blank=True))
        assert response.status_code==200,response.text
        assert requests==[[1]]
        document=response.json();assert 'OCR local' in document['notice']
        assert chat(client,conversation(client),'Leia o anexo.',[document['id']]).status_code==200
        supplied=local.contexts[-1][1][-1]['content']
        assert 'ALT-731' in supplied and 'ALT-732' in supplied and 'OCR local' in supplied


def test_document_excerpt_includes_its_notice_inside_the_budget():
    from backend.attachments import document_block
    row={'name':'long.txt','text':'Texto de teste. '*1000}
    block=document_block(row,1000,len)
    assert len(block)<=1000 and 'Trecho inicial' in block


def test_capabilities_come_from_the_selected_provider(tmp_path):
    _,local,_,client=fixture_client(tmp_path)
    async def capabilities(model):return ['completion','vision','tools']
    local.capabilities=capabilities
    with client:
        response=chat(client,conversation(client),'Voce consegue interpretar imagens e gerar imagens?')
        assert response.status_code==200,response.text
        supplied=local.contexts[-1][1][-1]['content']
        assert 'Interpretação de imagens: disponível' in supplied
        assert 'Geração de imagens: indisponível' in supplied
